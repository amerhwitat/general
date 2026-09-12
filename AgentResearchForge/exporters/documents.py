from pathlib import Path
from html import escape


def export_markdown(path: str, title: str, body: str, sources: list[dict] | None = None) -> str:
    target = Path(path); target.parent.mkdir(parents=True, exist_ok=True)
    text = f'# {title}\n\n{body}\n'
    if sources:
        text += '\n## Sources\n' + '\n'.join(f'- {s.get("title") or s.get("url")} — {s.get("url", "")}' for s in sources)
    target.write_text(text, encoding='utf-8')
    return str(target)


def export_html(path: str, title: str, body: str, sources: list[dict] | None = None) -> str:
    target = Path(path); target.parent.mkdir(parents=True, exist_ok=True)
    links = ''.join(f'<li><a href="{escape(s.get("url", ""))}">{escape(s.get("title") or s.get("url", "source"))}</a></li>' for s in (sources or []))
    html = f'<!doctype html><html><head><meta charset="utf-8"><title>{escape(title)}</title></head><body><h1>{escape(title)}</h1><p>{escape(body)}</p><h2>Sources</h2><ul>{links}</ul></body></html>'
    target.write_text(html, encoding='utf-8')
    return str(target)


def export_text(path: str, title: str, body: str) -> str:
    target = Path(path); target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(f'{title}\n\n{body}\n', encoding='utf-8')
    return str(target)


def export_docx(path: str, title: str, body: str, sources: list[dict] | None = None) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError('Install the documents extra for DOCX export') from exc
    target = Path(path); target.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(); doc.add_heading(title, 0)
    for paragraph in body.split('\n'):
        doc.add_paragraph(paragraph)
    if sources:
        doc.add_heading('Sources', level=1)
        for source in sources:
            doc.add_paragraph(f"{source.get('title') or source.get('url')} — {source.get('url', '')}")
    doc.save(target)
    return str(target)


def export_pdf(path: str, title: str, body: str, sources: list[dict] | None = None) -> str:
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
    except ImportError as exc:
        raise RuntimeError('Install the exports extra for PDF export') from exc
    target = Path(path); target.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(target), pagesize=letter)
    width, height = letter; y = height - 50
    c.setFont('Helvetica-Bold', 16); c.drawString(50, y, title); y -= 28
    c.setFont('Helvetica', 10)
    for line in (body + '\n\nSources:\n' + '\n'.join(s.get('url', '') for s in (sources or []))).splitlines():
        if y < 45: c.showPage(); y = height - 50; c.setFont('Helvetica', 10)
        c.drawString(50, y, line[:110]); y -= 14
    c.save(); return str(target)
