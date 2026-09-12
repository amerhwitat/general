from pathlib import Path
import re

SUPPORTED_TEXT = {'.txt','.md','.markdown','.html','.htm','.json','.csv','.xml','.yaml','.yml','.rst','.py','.js','.ts','.java','.c','.cpp','.h','.hpp','.rs','.go','.cs','.php','.rb','.swift','.kt','.kts','.sql','.sh','.ps1'}
OPTIONAL_DOCUMENTS = {'.pdf': 'pypdf', '.docx': 'docx', '.pptx': 'pptx', '.epub': 'ebooklib'}

class DocumentIndexer:
    def __init__(self):
        self.records: list[dict] = []

    def add_text(self, path: str, text: str, metadata: dict | None = None, chunk_size: int = 1200):
        clean = re.sub(r'\s+', ' ', text).strip()
        for start in range(0, len(clean), chunk_size):
            chunk = clean[start:start + chunk_size]
            if chunk:
                self.records.append({'path': path, 'text': chunk, 'metadata': metadata or {}, 'offset': start})

    def search(self, query: str, limit: int = 10) -> list[dict]:
        terms = set(re.findall(r'\w+', query.lower()))
        scored = []
        for record in self.records:
            hay = record['text'].lower()
            score = sum(hay.count(term) for term in terms)
            if score:
                scored.append((score, record))
        return [r for _, r in sorted(scored, key=lambda x: x[0], reverse=True)[:limit]]

    def ingest_tree(self, root: str, recursive: bool = True) -> int:
        base = Path(root)
        paths = base.rglob('*') if recursive else base.glob('*')
        count = 0
        for path in paths:
            if path.is_file():
                try:
                    count += int(self.ingest_path(str(path)))
                except (OSError, RuntimeError, ValueError):
                    continue
        return count

    def ingest_path(self, path: str):
        p = Path(path)
        if p.suffix.lower() in SUPPORTED_TEXT:
            self.add_text(str(p), p.read_text(encoding='utf-8', errors='replace'))
            return True
        suffix = p.suffix.lower()
        if suffix == '.pdf':
            try:
                from pypdf import PdfReader
            except ImportError as exc:
                raise RuntimeError('Install pypdf for PDF support') from exc
            reader = PdfReader(str(p))
            self.add_text(str(p), '\n'.join(page.extract_text() or '' for page in reader.pages))
            return True
        if suffix == '.docx':
            try:
                from docx import Document
            except ImportError as exc:
                raise RuntimeError('Install python-docx for DOCX support') from exc
            doc = Document(str(p))
            self.add_text(str(p), '\n'.join(x.text for x in doc.paragraphs))
            return True
        if suffix == '.pptx':
            try:
                from pptx import Presentation
            except ImportError as exc:
                raise RuntimeError('Install python-pptx for PPTX support') from exc
            deck = Presentation(str(p))
            text = '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, 'text'))
            self.add_text(str(p), text)
            return True
        return False
